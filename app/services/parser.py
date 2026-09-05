"""Extract text from uploaded PDF / PPTX / DOCX / TXT / MD files.

Each extracted page/section becomes a labeled chunk used for RAG and summaries.
"""
from __future__ import annotations

from pathlib import Path

# PDF -----------------------------------------------------------------


def _extract_pdf(path: Path) -> list[tuple[str, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    out: list[tuple[str, str]] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append((f"p. {i}", text))
    return out


# PPTX ----------------------------------------------------------------


def _extract_pptx(path: Path) -> list[tuple[str, str]]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[tuple[str, str]] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    parts.append(txt)
        if parts:
            out.append((f"slide {i}", "\n".join(parts)))
    return out


# DOCX ----------------------------------------------------------------


def _extract_docx(path: Path) -> list[tuple[str, str]]:
    import docx

    doc = docx.Document(str(path))
    out: list[tuple[str, str]] = []
    buf: list[str] = []
    page = 1
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        buf.append(text)
        if len(" ".join(buf)) > 1200:  # pseudo-pages keep chunks small
            out.append((f"section {page}", "\n".join(buf)))
            buf = []
            page += 1
    if buf:
        out.append((f"section {page}", "\n".join(buf)))
    return out


# TXT / MD ------------------------------------------------------------


def _extract_text(path: Path) -> list[tuple[str, str]]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    words, chunks, page = [], [], 1
    size = 0
    for para in raw.replace("\r\n", "\n").split("\n\n"):
        para = para.strip()
        if not para:
            continue
        words.append(para)
        size += len(para)
        if size > 1200:
            chunks.append("\n\n".join(words))
            words, size = [], 0
            page += 1
    if words:
        chunks.append("\n\n".join(words))
    return [(f"section {i}", c) for i, c in enumerate(chunks, start=1)]


def extract_chunks(path: Path, mime_type: str = "") -> list[tuple[str, str]]:
    """Return a list of (label, text) chunks for the given file."""
    suffix = path.suffix.lower()
    if suffix == ".pdf" or "pdf" in mime_type:
        return _extract_pdf(path)
    if suffix == ".pptx" or "presentation" in mime_type:
        return _extract_pptx(path)
    if suffix == ".docx" or "wordprocessingml" in mime_type:
        return _extract_docx(path)
    return _extract_text(path)
